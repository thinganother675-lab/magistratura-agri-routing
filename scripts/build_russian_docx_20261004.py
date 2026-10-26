"""Editable dated course draft; run with the bundled artifact Python."""
from pathlib import Path
import json
import re
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript/russian_resource_20261004.md"
TARGET = SOURCE.with_suffix(".docx")


def add_text(paragraph, text, font_size=14):
    # Preserve source words and link destinations; no generated factual content.
    for part in re.split(r"(\[[^\]]+\]\(https?://[^\s)]+\)|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^\s)]+)\)", part)
        if link:
            relation = paragraph.part.relate_to(link[2], RT.HYPERLINK, is_external=True)
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), relation)
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            size = OxmlElement("w:sz")
            size.set(qn("w:val"), str(2 * font_size))
            props.append(size)
            run.append(props)
            content = OxmlElement("w:t")
            content.text = link[1]
            run.append(content)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
            continue
        run = paragraph.add_run(part.strip("*`") if part.startswith(("*", "`")) else part)
        run.font.size = Pt(font_size)
        if part.startswith("**"):
            run.bold = True
        elif part.startswith("*"):
            run.italic = True


def set_font(style, size):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), "Times New Roman")
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)


def main():
    d = Document()
    section = d.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    section.footer_distance = Cm(1)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        set_font(d.styles[name], 14)
        f = d.styles[name].paragraph_format
        f.space_before = f.space_after = Pt(0)
        f.line_spacing = 1
        f.first_line_indent = Cm(0 if name != "Normal" else 1.25)
        f.widow_control = True
    for border in d.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for name in ("Title", "Heading 1", "Heading 2"):
        d.styles[name].font.bold = True
        d.styles[name].paragraph_format.keep_with_next = True
    d.styles["Title"].paragraph_format.space_after = Pt(10)
    d.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    d.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    d.styles["Heading 2"].paragraph_format.space_before = Pt(8)
    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    foot._p.append(field)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_refs = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            p = d.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(p, line[2:])
        elif line.startswith("## "):
            in_refs = line == "## Литература"
            d.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            d.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                row = [x.strip() for x in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r"[:\-]+", x) for x in row):
                    rows.append(row)
                i += 1
            n = len(rows[0])
            widths = {3: [8, 4.5, 4.5], 4: [6, 4.5, 2.2, 4.3], 5: [5, 3, 3, 3, 3]}[n]
            table = d.add_table(rows=len(rows), cols=n)
            table.autofit = False
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            props = table._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement("w:" + edge)
                for k, v in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    el.set(qn("w:" + k), v)
                borders.append(el)
            props.append(borders)
            for j, row in enumerate(rows):
                if j == 0:
                    table.rows[j]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                table.rows[j]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for k, value in enumerate(row):
                    c = table.cell(j, k)
                    c.width = Cm(widths[k])
                    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = c.paragraphs[0]
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.keep_with_next = j < len(rows) - 1
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if k == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    add_text(p, value)
                    for run in p.runs:
                        run.font.size = Pt(12)
                        run.bold = j == 0
                    if j == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "EDEDED")
                        c._tc.get_or_add_tcPr().append(shade)
            continue
        else:
            p = d.add_paragraph()
            add_text(p, line, 12 if in_refs or line.startswith("Таблица ") else 14)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if line.startswith("Таблица "):
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_before = Pt(5)
            if in_refs:
                p.paragraph_format.first_line_indent = Cm(-0.7)
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(12)
            elif i <= 7:
                p.paragraph_format.first_line_indent = Cm(0)
        i += 1
    d.core_properties.title = lines[0][2:]
    d.core_properties.author = "Антоний Сарычев; Денис Шестернин"
    d.core_properties.subject = "Исследовательский черновик; согласие авторов и лицензия ожидаются"
    d.core_properties.language = "ru-RU"
    d.save(TARGET)
    # The editable source contains plain-text simple formulas, not raster equations.
    print(json.dumps({"path": str(TARGET), "created_utc": datetime.now(timezone.utc).isoformat(), "content_source": str(SOURCE), "visual_review": "PENDING"}, ensure_ascii=False))


if __name__ == "__main__":
    main()

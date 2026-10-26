from pathlib import Path
from datetime import datetime,timezone,timedelta
import json,hashlib,re,textwrap,math,xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
R=Path(__file__).resolve().parent; C=R.parent
for f in ['submission','evidence','evidence/diagrams','qa']:(R/f).mkdir(parents=True,exist_ok=True)
def rows(t):return [[v.strip() for v in line.split('|')] for line in t.strip().splitlines()]
D={
1:dict(title='Паспорт выдачи кормов',diagram='context',tables=[rows('''Поле|Содержательное заполнение
Объект и граница|Выдача кормов со склада С01 группам G01/G02 и подготовка суточного DOC-04 для директора Р04. Синтетический кейс К01–К10.
Проблема|APP-03 суммирует числовые строки без автоматического контроля OP, статуса и полноты; повтор пересылки может увеличить итог.
Кому мешает|Р03 вручную сопоставляет выдачи и подтверждения; Р04 получает сумму без объяснения нерешённых случаев.
Вопрос проекта|Как обеспечить однократный учёт подтверждённых выдач и явно показывать ограничения суточного итога?
Output / outcome|Output: согласованные правила и описание учёта. Outcome: ожидаемые сокращение ручной сверки и повышение достоверности; пока гипотеза, а не измеренный эффект.
Вне границы|Рационы, ветеринария, закупки, бухгалтерские проводки и замена APP-04.'''),rows('''Утверждение|Основание|Статус
Выдачи и подтверждения поступают разными путями, Р03 сопоставляет их вручную.|К03|Факт учебного кейса.
Единый шаблон с контролем OP и регламентом приёмки может уменьшить повторный учёт.|К02–К03; причинный вывод для пути A|Проектная гипотеза; эффект не измерен.
Точная задержка каждого сообщения и частота ошибочных отчётов за год.|В К01–К10 значений нет.|Неизвестно; числа не подставляются.''')],
diagnosis='Критерий — различить физический результат и качество сведений о нём. Корм поступает группам, а директор получает DOC-04. По К03 раздельные источники и повторная пересылка при суммировании строк скрывают повтор и неполноту. Для пути A предлагается контролируемый шаблон с устойчивым OP и отдельным списком нерешённых случаев; само наличие файла ещё не доказывает соблюдение правил. В дальнейшей проверке две доставки одной выдачи должны дать один вклад, а нерешённый OP — запрет статуса «полный».',
newcase='Для Р02 в DOC-04 добавляем group_code и confirmed_quantity_kg — подтверждённую массу отдельно по каждой группе. Первое поле определяет объект разбора, второе показывает, сколько по нему достоверно выдано. Для оценки отклонения Р02 сопоставляет эту массу с потребностью DOC-01; конкретные плановые количества в кейсе не заданы и здесь не выдумываются. Список нерешённых OP и статус полноты сохраняются, чтобы неполные сведения не выглядели точным расходом.',
check='Агент сверил К01–К04: складская выдача на схеме остаётся материальным потоком, DOC-04 — информационным. Прогноз пользы отделён от наблюдения; личная проверка Денисом не заявляется.',
sources='К01–К04, К08–К09; PZ01_Instruction.docx. Л1: 65–79, 95–99; Л3: 35–36. API APP-04 и годовая статистика ошибок не известны.'),
2:dict(title='Домены учёта кормов',diagram='domains',tables=[rows('''Домен|Объекты AS-IS|Значимая связь
Бизнес|Выдача Р01, подтверждение Р02, сверка Р03 и потребление отчёта Р04.|Полномочие Р02 определяет, чьё решение подтверждает хозяйственный смысл записи.
Данные|D01 выдача, D04 подтверждение; DOC-02, DOC-03 и DOC-04.|OP должен связывать факт выдачи с подтверждением; код доставки M не заменяет OP.
Приложения|APP-01 складской файл, APP-02 сообщения, APP-03 свод; APP-04 внешний контур.|Р03 переносит строки и вручную сопоставляет подтверждения; APP-03 автоматически складывает числовые строки.
Технологии|T01 ноутбук склада, T02 место бухгалтера, T03 файловый сервер, T04 связь.|Локальный ввод на T01 продолжается при отказе T04, но общий свод получает сведения позже.'''),rows('''Роль|Вопрос|Что показать
Р01|Какая запись выдачи передана и принята?|Реестр передач с OP и состоянием доставки; для проекта, не существующий журнал AS-IS.
Р02|Какую выдачу и количество требуется подтвердить?|Карточку OP с кормом, группой, массой и ролью автора.
Р03|Какие строки можно включать в итог?|Таблицу правил сопоставления и представление нерешённых OP.
Р04|Какая часть суточного отчёта достоверна?|DOC-04: итог, полнота, учтённые и ожидаемые OP.
Р05|От каких узлов зависит доступность файлов?|Схему размещения приложений, канала и копии с границами неизвестного.''')],
diagnosis='Схема показывает две зависимости AS-IS: получение DOC-04 → OP/статус → ручная сверка в APP-03 → T02/T03; подтверждение Р02 → DOC-03 → APP-02 → устройства Р02/Р03 и канал. К03 не задаёт автоматического списка исключений. Для пути A его добавление меняет одновременно содержание DOC-04, контролируемый шаблон и работу Р03. API APP-04 неизвестен; вопрос владельцу: какой проверяемый формат согласованного итога система принимает?',
newcase='Добавления OP в DOC-03 недостаточно, пока APP-03 связывает записи только по названию корма. В пути A нужно изменить правило сверки и ссылку в шаблоне на operation_id, а корм и группу оставить проверяемыми атрибутами. При отсутствии однозначного совпадения Р03 должен видеть исключение и направить его Р02. Приёмочный пример — две разные выдачи одного корма: OP должны остаться разными хозяйственными фактами.',
check='Агент сверил две цепочки по К03–К05: ручное сопоставление не обозначено API или автоматической дедупликацией. Реестр приёмки — проектное представление для Р01, а не выдуманный факт текущего состояния.',
sources='К01–К05, К08–К09; ПЗ-1; PZ02_Instruction.docx. Л2: 22–40, 90–95. Устройства и API APP-04 не обследованы.'),
3:dict(title='Цели достоверного отчёта',diagram='goals',tables=[rows('''Элемент|Формулировка|Почему это не средство
Общая цель|Р04 получает объяснимый суточный отчёт о подтверждённых выдачах.|Определён результат деятельности; продукт не назван целью.
Подцель 1|Одна выдача OP даёт ровно один вклад при точном повторе доставки.|Проверяемое свойство учёта.
Подцель 2|Нерешённые OP и неполнота DOC-04 видимы потребителю.|Свойство результата, а не покупка системы.
Подцель 3|Активная подготовка сопоставимого отчёта занимает не более 40 минут.|Ориентир К08; 40 минут пути A по К09 — прогноз, не факт.
Функция 1|Регистрация физической выдачи: Р01 записывает DOC-02 в APP-01.|Устойчивая деятельность; технология её реализации может меняться.
Функция 2|Подтверждение производственного смысла: Р02 формирует DOC-03.|Хозяйственное полномочие, не техническая доставка.
Функция 3|Контроль сведений и свод: Р03 сопоставляет DOC-02/DOC-03 и готовит DOC-04.|Часть процесса К03; не название приложения.'''),rows('''Показатель|AS-IS и основание|Цель и проверка
Активное время|К07: (75+90+80+70+85)/5 = 400/5 = 80 минут после получения файлов.|≤40 минут при одинаковых входах, правилах полноты и границе замера; будущий замер до/после.
Повторный вклад|К03: нет автоматического контроля OP; годовая частота ошибок неизвестна.|Две доставки одного неизменяемого OP не увеличивают итог; проверить на модели позже.
Полнота отчёта|К03: статус автоматически не рассчитывается.|Сопоставить ожидаемые и разрешённые OP; при нерешённых показать неполный статус и перечень причин.''')],
diagnosis='Среднее 80 минут относится к заданным активным работам, а не ко всей смене. Разложение 15+20+35+10=80 из К07 описывает типовой состав, а не дополнительные пять измерений. В пути A контролируемый шаблон должен поддерживать OP, состояния и перечень исключений, а регламент — приёмку подтверждений Р02 и сверку Р03. Поэтому цель полноты меняет и данные DOC-04, и функцию подготовки свода; покупка нового сервиса не является необходимым условием цели. Фактическое достижение 40 минут пока не доказано.',
newcase='Сделать вывод о более раннем поступлении отчёта нельзя: 20 минут сверки и два часа ожидания относятся к разным границам времени. Добавляем показатель «календарная задержка» от согласованного начала цикла до получения DOC-04 директором. В нём отдельно учитываем ожидание подтверждения и остальные действия процесса; 20 минут нельзя автоматически считать временем всей подготовки. Сравнение выполняем при одинаковом составе входов и той же полноте, а исходная сопоставимая задержка в кейсе не задана.',
check='Агент пересчитал учебную выборку: сумма 400, среднее 80 минут; отдельно проверено разложение 15+20+35+10. Активное время, ожидание и прогноз A не смешаны; личные измерения Дениса и фактический результат внедрения не заявлены.',
sources='К01–К03, К07–К09; ПЗ-1–2; PZ03_Instruction.docx. Л3: 4–18, 46–51, 87–88; Л1: 77–78. Частота ошибок и исходная календарная задержка неизвестны.')
}

# Compact wording keeps all required results together on readable pages.
D[1]['diagnosis']='Критерий — различить физический результат и сведения о нём. Корм поступает группам, DOC-04 — директору. По К03 повторная пересылка и суммирование строк могут скрыть повтор и неполноту. Для пути A предлагается контролируемый шаблон с OP и списком нерешённых случаев. Проверка позже должна показать один вклад при повторе и запрет статуса «полный» при нерешённом OP; существующий файл этого ещё не гарантирует.'
D[1]['newcase']='В DOC-04 для Р02 добавляем group_code и confirmed_quantity_kg — подтверждённую массу по группе. Первое поле определяет объект разбора, второе показывает достоверно выданное количество. Для оценки отклонения Р02 сравнивает массу с потребностью DOC-01; плановые количества в кейсе не заданы и не выдумываются. Статус полноты и нерешённые OP сохраняются, чтобы неполные сведения не выглядели точным расходом.'
D[1]['check']='Агент сверил К01–К04: корм — материальный выход, DOC-04 — информационный. Польза шаблона оставлена гипотезой; личная проверка Денисом не заявляется.'
D[2]['diagnosis']='Обе цепочки AS-IS подписаны на схеме. По К03 APP-03 не формирует автоматический список исключений. В пути A его добавление меняет DOC-04, контролируемый шаблон и работу Р03 одновременно. Неизвестный API APP-04 не принят за существующий: владельцу нужен вопрос о доступном формате согласованного итога. Новые представления для участников — требования к ответу, а не доказательство их наличия сейчас.'
D[2]['newcase']='Одного OP в DOC-03 недостаточно, пока APP-03 связывает записи только по названию корма. В пути A меняем правило сверки и ссылку в шаблоне на operation_id; корм и группу проверяем как атрибуты. Неоднозначное совпадение Р03 передаёт Р02 как исключение. Приёмочный пример: две выдачи одного корма с разными OP должны остаться разными фактами.'
D[2]['check']='Агент сверил две цепочки по К03–К05: ручное сопоставление не названо API. Реестр приёмки Р01 обозначен проектным представлением, а не фактом AS-IS.'
D[3]['diagnosis']='Среднее 80 минут относится к активной работе после получения файлов. Разложение 15+20+35+10=80 из К07 — состав типового процесса, а не дополнительные измерения. В пути A полнота требует OP, состояний и исключений в шаблоне, а также регламента подтверждения Р02 и сверки Р03. Поэтому меняются данные DOC-04 и функция свода. Достижение 40 минут пока не доказано.'
D[3]['newcase']='Более раннее поступление отчёта не доказано: 20 минут сверки и два часа ожидания имеют разные границы. Добавляем показатель календарной задержки от согласованного начала цикла до получения DOC-04 директором. В нём отдельно учитываем ожидание и остальные действия, поэтому 20 минут не принимаем за всю подготовку. Сравнение требует одинаковых входов и полноты; исходная сопоставимая задержка в кейсе неизвестна.'
D[3]['check']='Агент пересчитал выборку: 400/5=80 минут, отдельно — 15+20+35+10=80. Активное время, ожидание и прогноз A различены; личные измерения Дениса не заявлены.'
D[3]['tables'][0]=D[3]['tables'][0][:5]+[['Три функции','1. Р01 регистрирует выдачу в DOC-02/APP-01; 2. Р02 подтверждает смысл в DOC-03; 3. Р03 контролирует сведения и формирует DOC-04.','Устойчивые виды деятельности в К03; их реализация может меняться.']]
graphs={
'context':('AS-IS   Выдача со склада и получение отчёта',[
 ('need','Потребность Р02\nDOC-01',70,130),('issue','Выдача Р01\nсклад С01',590,130),('group','G01 / G02\nполучают корм',1110,130),('info','Сведения DOC-02\nподтверждение DOC-03',70,510),('report','Р03 готовит\nотчёт DOC-04',590,510),('director','Директор Р04\nполучает итог',1110,510)],
 [('need','issue','данные'),('issue','group','корм, кг'),('issue','info','сведения о факте'),('info','report','данные'),('report','director','DOC-04')]),
'domains':('AS-IS   Две цепочки зависимостей',[
 ('b1','Бизнес\nполучение DOC-04',140,110),('d1','Данные\nOP и статус',140,330),('a1','APP-03\nручная сверка и сумма',140,550),('t1','T02 / T03\nрабочее место и файлы',140,770),
 ('b2','Бизнес\nподтверждение Р02',980,110),('d2','DOC-03\nOP и решение Р02',980,330),('a2','APP-02\nканал сообщений',980,550),('t2','Устройства Р02/Р03\nи канал связи',980,770)],
 [('b1','d1','требует'),('d1','a1','используются'),('a1','t1','зависит от'),('b2','d2','создаёт'),('d2','a2','передаётся'),('a2','t2','зависит от')]),
'goals':('Проектные цели   Контроль важнее ускорения',[
 ('g','Объяснимый отчёт\nдля Р04',590,110),('one','Один OP\nодин вклад',70,480),('full','Неполнота и причины\nвидимы в DOC-04',590,480),('time','Активная подготовка\nне более 40 минут',1110,480)],
 [('g','one','подцель 1'),('g','full','подцель 2'),('g','time','подцель 3')])}
def diagram(name,graph):
    title,nodes,edges=graph;w=1600;h=max(y for _,_,_,y in nodes)+205;bw=390;bh=140
    im=Image.new('RGB',(w,h),'white');draw=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',30);label_font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',25)
    draw.text((35,28),title,font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',34),fill='#18342e')
    positions={n:(x,y) for n,_,x,y in nodes}
    for a,b,label in edges:
        ax,ay=positions[a];bx,by=positions[b];dx=bx-ax;dy=by-ay
        ac=(ax+bw/2,ay+bh/2);bc=(bx+bw/2,by+bh/2);scale=max(abs(dx)/(bw/2),abs(dy)/(bh/2))
        sx,sy=ac[0]+dx/scale,ac[1]+dy/scale;ex,ey=bc[0]-dx/scale,bc[1]-dy/scale
        draw.line((sx,sy,ex,ey),fill='#24564b',width=3)
        angle=math.atan2(ey-sy,ex-sx)
        draw.polygon([(ex,ey),(ex-17*math.cos(angle-.45),ey-17*math.sin(angle-.45)),(ex-17*math.cos(angle+.45),ey-17*math.sin(angle+.45))],fill='#24564b')
        mid=((sx+ex)/2,(sy+ey)/2);wrapped='\n'.join(textwrap.wrap(label,15))
        box=draw.multiline_textbbox((0,0),wrapped,font=label_font);tw=box[2];th=box[3]
        draw.rectangle((mid[0]-tw/2-5,mid[1]-th/2-6,mid[0]+tw/2+5,mid[1]+th/2+6),fill='white')
        draw.multiline_text(mid,wrapped,font=label_font,fill='black',anchor='mm',align='center')
    for n,label,x,y in nodes:
        draw.rounded_rectangle((x,y,x+bw,y+bh),radius=10,outline='#24564b',width=3,fill='#eef4f0')
        draw.multiline_text((x+bw/2,y+bh/2),label,font=font,fill='black',anchor='mm',align='center',spacing=7)
    im.save(R/'evidence/diagrams'/f'{name}.png')
    root=ET.Element('mxfile');dia=ET.SubElement(root,'diagram',name=name);model=ET.SubElement(dia,'mxGraphModel');cells=ET.SubElement(model,'root');ET.SubElement(cells,'mxCell',id='0');ET.SubElement(cells,'mxCell',id='1',parent='0')
    for n,label,x,y in nodes:
        cell=ET.SubElement(cells,'mxCell',id=n,value=label,vertex='1',parent='1',style='rounded=1;whiteSpace=wrap;html=0;fontSize=26;fillColor=#eef4f0;strokeColor=#24564b;')
        ET.SubElement(cell,'mxGeometry',x=str(x),y=str(y),width=str(bw),height=str(bh),**{'as':'geometry'})
    for i,(a,b,label) in enumerate(edges):
        cell=ET.SubElement(cells,'mxCell',id=f'e{i}',source=a,target=b,value=label,edge='1',parent='1',style='endArrow=block;fontSize=22;')
        ET.SubElement(cell,'mxGeometry',relative='1',**{'as':'geometry'})
    ET.ElementTree(root).write(R/'evidence/diagrams'/f'{name}.drawio',encoding='utf-8',xml_declaration=True)
for name,graph in graphs.items():diagram(name,graph)

def replace(p,t):
    if p.runs:p.runs[0].text=t
    else:p.add_run(t)
    for run in p.runs[1:]:run.text=''
def fill(table,data):
    while len(table.rows)<len(data):table.add_row()
    while len(table.rows)>len(data):table._tbl.remove(table.rows[-1]._tr)
    for i,values in enumerate(data):
        for j,value in enumerate(values):
            cell=table.cell(i,j);replace(cell.paragraphs[0],value)
            for p in cell.paragraphs[1:]:p._element.getparent().remove(p._element)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(3)
                for r in p.runs:r.font.size=Pt(11.5)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D2DDD7');borders.append(el)
            cell._tc.get_or_add_tcPr().append(borders)
    for row in table.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
templates=[]
for n,data in D.items():
    source=C/'originals'/f'Answer_PZ{n:02}.docx';doc=Document(source)
    templates.append(dict(pz=n,source=str(source),sha256=hashlib.sha256(source.read_bytes()).hexdigest(),size=[doc.sections[0].page_width.cm,doc.sections[0].page_height.cm]))
    assert len(doc.tables)==2
    widths={1:[[4,12],[8,3.5,4.5]],2:[[2.6,6.7,6.7],[1.6,6.4,8]],3:[[2.4,8,5.6],[3,6,7]]}
    for ti,(table,values) in enumerate(zip(doc.tables,data['tables'])):
        fill(table,values);table.autofit=False
        for j,width in enumerate(widths[n][ti]):
            table.columns[j].width=Cm(width)
            for cell in table.columns[j].cells:cell.width=Cm(width)
    # Keep the source form structure and A4 sections, remove its decorative title rule.
    title_style=doc.styles['Title'];title_style.font.color.rgb=RGBColor(0,0,0);title_style.font.underline=False
    for b in list(title_style.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
    for p in list(doc.paragraphs):
        t=p.text
        if t.startswith(f'ПЗ-{n:02}. Бланк'):replace(p,data['title'])
        elif t.startswith('ФИО:'):replace(p,f'Шестернин Денис · ДЭ 39-26 · ПЗ-{n:02} · 16 октября 2026')
        elif t.startswith('Объект:'):replace(p,'Объект: учебный кейс К01–К10. Направление A: контролируемые файлы и регламент.')
        elif t.startswith('Заполняйте результат'):replace(p,'Тестовый образец подготовлен агентом Codex. Личная проверка Денисом не заявляется.')
        elif t=='Моя схема и правила чтения':replace(p,'Схема и правила чтения')
        elif t.startswith('[Вставьте схему'):
            replace(p,'');p.add_run().add_picture(str(R/'evidence/diagrams'/f"{data['diagram']}.png"),width=Cm(16))
        elif t.startswith('Критерий: ______'):replace(p,data['diagnosis'])
        elif t.startswith('[Назовите конкретное'):replace(p,data['newcase'])
        elif t=='Что я лично проверил':replace(p,'Применение ИИ и проверка агента')
        elif t.startswith('[Один конкретный'):replace(p,data['check'])
        elif t.startswith('[Укажите К-разделы'):replace(p,data['sources']);[setattr(r.font,'size',Pt(11)) for r in p.runs]
    if n==3:
        # The third form has no diagram slot; add the required goal tree before diagnosis.
        heading=next(p for p in doc.paragraphs if p.text=='Диагноз и междоменные последствия')
        cap=heading.insert_paragraph_before('Дерево целей',style='Heading 2')
        image_p=heading.insert_paragraph_before();image_p.add_run().add_picture(str(R/'evidence/diagrams/goals.png'),width=Cm(16))
    for p in list(doc.paragraphs):
        if p.text=='Ответ на новое условие':
            nxt=p._element.getnext()
            assert nxt is not None and nxt.tag==qn('w:p')
            nxt.getparent().remove(nxt)
        if p.style.name=='Title':
            for b in list(p._element.iter(qn('w:pBdr'))):b.getparent().remove(b)
            for r in p.runs:r.font.color.rgb=RGBColor(0,0,0);r.font.underline=False
            p.paragraph_format.space_after=Pt(5)
        if p.style.name.startswith('Heading'):
            p.paragraph_format.keep_with_next=True
            for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
        if p.style.name=='Normal':p.paragraph_format.space_after=Pt(4)
        if p.text in (data['newcase'],data['check']):p.paragraph_format.keep_together=True
    dest=R/'submission'/f'PZ{n:02}_Шестернин_ДЭ39-26.docx';doc.save(dest)
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert not re.search(r'______|\[Вставьте|\[Ваш|\[Назовите',text)
    assert 'Сарычев' not in text and 'Шестернин Денис' in text
    (R/'evidence'/f'PZ{n:02}_text.txt').write_text(text,encoding='utf-8')
(R/'evidence/template_evidence.json').write_text(json.dumps(templates,ensure_ascii=False,indent=2),encoding='utf-8')
(R/'evidence/content.json').write_text(json.dumps(D,ensure_ascii=False,indent=2),encoding='utf-8')
(R/'evidence/artifact.md').write_text('Три тестовые работы Дениса по Answer_PZ01–03.docx. Сохранены исходная структура таблиц, секции A4, шрифты и иерархия. Заполнены предусмотренные поля, добавлены строки функций и дерево целей ПЗ-3. Текст таблиц 11.5 pt, основной текст исходного бланка 14 pt. Декоративная линия под заголовком удалена. ФИО без неподтверждённого отчества; группа из прежних материалов. Личные действия Денису не приписаны. Источники и состояние AS-IS отделены от предложений пути A. Обязательна проверка рендера Word/PDFium.',encoding='utf-8')
print('Created 3 DOCX for Denis, variant A; no submissions.')

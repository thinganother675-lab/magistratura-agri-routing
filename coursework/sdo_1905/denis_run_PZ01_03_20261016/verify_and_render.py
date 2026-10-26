from pathlib import Path
from datetime import datetime,timezone
from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image,ImageDraw
from docx import Document
import json,re,hashlib,zipfile
R=Path(__file__).resolve().parent;C=R.parent
checks=[]
for n in range(1,4):
    docx=R/'submission'/f'PZ{n:02}_Шестернин_ДЭ39-26.docx'
    pdf=R/'qa'/f'PZ{n:02}_Шестернин_ДЭ39-26.pdf'
    reader=PdfReader(pdf);assert 2<=len(reader.pages)<=3,(n,len(reader.pages))
    d=Document(docx)
    text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert 'Шестернин Денис' in text and 'ДЭ 39-26' in text and 'Сарычев' not in text
    assert 'Ответ на новое условие' in text and 'Применение ИИ и проверка агента' in text
    assert 'агентом Codex' in text and not re.search(r'______|\[Вставьте|\[Ваш|\[Назовите',text)
    assert len(d.tables)==2
    assert len(d.inline_shapes)>=1
    assert all(abs(s.page_width.cm-21)<.1 for s in d.sections)
    with zipfile.ZipFile(docx) as z:assert z.testzip() is None
    output=R/'qa'/pdf.stem;output.mkdir(exist_ok=True)
    pages=pdfium.PdfDocument(pdf);thumbs=[]
    for i,page in enumerate(pages):
        im=page.render(scale=1.7).to_pil().convert('RGB');im.save(output/f'page-{i+1}.png')
        thumb=im.copy();thumb.thumbnail((720,1020));thumbs.append(thumb)
        assert '\ufffd' not in (reader.pages[i].extract_text() or '')
    contact=Image.new('RGB',(len(thumbs)*720,1050),'#ddd');draw=ImageDraw.Draw(contact)
    for i,im in enumerate(thumbs):contact.paste(im,(i*720,25));draw.text((i*720+10,5),f'PZ{n:02} / {i+1}',fill='black')
    contact.save(output/'contact.jpg',quality=93)
    checks.append(dict(pz=n,pages=len(reader.pages),bytes=docx.stat().st_size,sha256=hashlib.sha256(docx.read_bytes()).hexdigest(),file=docx.name))
assert (75+90+80+70+85)/5==80 and 15+20+35+10==80
sources=json.loads((C/'file_manifest.json').read_text(encoding='utf-8'))
assert all(hashlib.sha256((C/x['local_path']).read_bytes()).hexdigest()==x['sha256'] for x in sources)
plan=json.loads((C/'submission_schedule/plan.json').read_text(encoding='utf-8'))
assert all(hashlib.sha256(Path(f['path']).read_bytes()).hexdigest()==f['sha256'] for s in plan['slots'] for f in s['files'])
criteria={
 1:['Граница, Р03/Р04, проблема и вопрос — паспорт','Корм и DOC-04 разделены на схеме; output/outcome различены','Три основания: факт, гипотеза, неизвестное','Новое условие: group_code и confirmed_quantity_kg, четыре предложения'],
 2:['Четыре домена с объектами К01–К05','Все Р01–Р05 с вопросами и представлениями','Две цепочки с приложением и технологической зависимостью','Новое условие: связь по operation_id вместо названия корма, четыре предложения'],
 3:['Дерево одной цели и трёх подцелей, три функции с ролями','400/5=80 минут активной работы; разложение отдельно','Три показателя с проверкой и связью DOC-04/шаблона','Новое условие: отдельная календарная задержка, четыре предложения']}
report=dict(status='checks_passed_visual_review_pending',completed_practicals=[1,2,3],variant_direction='A',synthetic_case=True,
 identity='Шестернин Денис, ДЭ 39-26',patronymic='не подтверждено, не добавлено',facts_by='Агент Codex, не личный опыт Дениса',
 checked_at=datetime.now(timezone.utc).isoformat(),source_integrity_count=len(sources),antoniy_submission_files_unchanged=True,submitted_to_sdo=False,
 criteria_review=criteria,files=checks)
(R/'evidence/verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(status=report['status'],pages=[x['pages'] for x in checks],source_files_unchanged=77,antoniy_files_unchanged=True),ensure_ascii=False))

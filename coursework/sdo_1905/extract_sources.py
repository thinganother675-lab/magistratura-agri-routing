from pathlib import Path
import json, hashlib, zipfile
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

ROOT=Path(__file__).parent
records=[]
for p in sorted((ROOT/'originals').iterdir()):
    record={'name':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
    text=''
    if p.suffix=='.pdf':
        pdf=PdfReader(p)
        record['pages']=len(pdf.pages)
        text='\n\n'.join(f'--- Страница {i+1} ---\n'+(page.extract_text() or '') for i,page in enumerate(pdf.pages))
    elif p.suffix=='.docx':
        doc=Document(p)
        parts=[]
        for node in doc.element.body:
            if node.tag.endswith('}p'):
                from docx.text.paragraph import Paragraph
                parts.append(Paragraph(node,doc).text)
            elif node.tag.endswith('}tbl'):
                from docx.table import Table
                for row in Table(node,doc).rows:
                    parts.append(' | '.join(c.text for c in row.cells))
        text='\n'.join(parts)
        with zipfile.ZipFile(p) as z:
            record['embedded_media']=[n for n in z.namelist() if n.startswith('word/media/')]
    elif p.suffix in ('.txt','.csv'):
        text=p.read_text(encoding='utf-8-sig')
    elif p.suffix=='.xlsx':
        wb=load_workbook(p,data_only=False)
        sheets=[]
        for s in wb:
            sheets.append(f'## {s.title} ({s.max_row} rows x {s.max_column} cols)')
            sheets.extend(' | '.join('' if c.value is None else f'{c.coordinate}: {c.value}' for c in row) for row in s if any(c.value is not None for c in row))
        text='\n'.join(sheets)
        record['sheets']=wb.sheetnames
    elif p.suffix=='.drawio':
        import xml.etree.ElementTree as ET
        tree=ET.parse(p)
        record['diagrams']=[e.attrib.get('name') for e in tree.findall('.//diagram')]
        text=p.read_text(encoding='utf-8')
    if text:
        (ROOT/'extracted'/f'{p.name}.txt').write_text(text,encoding='utf-8')
        record['text_chars']=len(text)
    records.append(record)
(ROOT/'file_manifest.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'files':len(records),'pdf_pages':sum(r.get('pages',0) for r in records),'bytes':sum(r['bytes'] for r in records)},ensure_ascii=False))

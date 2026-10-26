"""Export the same checked report content to editable Word paragraphs and tables."""
import re
from html.parser import HTMLParser
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.platypus import Paragraph, Table, Image, PageBreak, Spacer

class Inline(HTMLParser):
    def __init__(self, paragraph):
        super().__init__(convert_charrefs=True); self.p=paragraph; self.tags=[];self.href=None
    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        if tag=='link':self.href=dict(attrs).get('href')
        if tag=='br': self.p.add_run().add_break()
    def handle_endtag(self, tag):
        if tag in self.tags: self.tags.remove(tag)
        if tag=='link':self.href=None
    def handle_data(self, data):
        run=self.p.add_run(data)
        run.bold=True if 'b' in self.tags or 'strong' in self.tags else None
        run.italic='i' in self.tags
        run.font.subscript='sub' in self.tags
        run.font.superscript='super' in self.tags or 'sup' in self.tags
        if self.href:
            hyperlink=OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'),self.p.part.relate_to(self.href,RT.HYPERLINK,is_external=True))
            hyperlink.append(run._r);self.p._p.append(hyperlink)

def paragraph(container, item, cell=False):
    name=item.style.name
    text=item.text
    if name in ['TitleRU','H1RU']:
        text=re.sub(r'^\d+\.\s*','',text).replace('№','').replace(':',' ').replace('.',' ').replace('—',' ').replace('-',' ').strip()
    p=container.add_paragraph(style='Title' if name=='TitleRU' else ('Heading 1' if name=='H1RU' else 'Normal'))
    Inline(p).feed(text)
    if name in ['SmallRU','CellRU','TinyRU','CodeRU'] or cell:
        for run in p.runs:
            run.font.size=Pt(7.1 if name=='TinyRU' else (8.2 if cell else (8 if name=='CodeRU' else 9)))
            if name=='CodeRU':run.font.name='Consolas'
        p.paragraph_format.space_after=Pt(3 if cell else 5)
        p.paragraph_format.line_spacing=1.05
    p.paragraph_format.keep_with_next=name in ['TitleRU','H1RU']
    return p

def export_docx(story, path, person, variant):
    doc=Document(); section=doc.sections[0]
    # The bundled base template includes decorative paragraph borders.
    # Clear them from styles so the formal report has plain black titles.
    for style in doc.styles:
        for border in list(style.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    section.page_width=Pt(595.276);section.page_height=Pt(841.89)
    section.left_margin=section.right_margin=Pt(45)
    section.top_margin=Pt(40);section.bottom_margin=Pt(55)
    section.footer_distance=Pt(22)
    for name,size in [('Normal',11),('Title',18),('Heading 1',14)]:
        s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_after=Pt(8)
        s.paragraph_format.line_spacing=1.10
        s.paragraph_format.widow_control=True
        if name!='Normal':s.font.bold=True
    doc.core_properties.author=person
    doc.core_properties.title=f'Лабораторная работа вариант {variant}'
    foot=section.footer.paragraphs[0]
    foot.add_run(f'Вариант {variant}   ДЭ 39-26   GAMS 54.5.0').font.size=Pt(8)
    foot.add_run(' '*12)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    for item in story:
        if isinstance(item,Paragraph):paragraph(doc,item)
        elif isinstance(item,PageBreak):doc.add_page_break()
        elif isinstance(item,Spacer):
            # Paragraph spacing already preserves vertical rhythm.
            continue
        elif isinstance(item,Image):
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
            p.add_run().add_picture(str(item.filename),width=Pt(item.drawWidth),height=Pt(item.drawHeight))
        elif isinstance(item,Table):
            data=item._cellvalues;table=doc.add_table(rows=len(data),cols=len(data[0]));table.autofit=False
            widths=item._argW
            for rowidx,(row,values) in enumerate(zip(table.rows,data)):
                trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
                if rowidx==0:
                    repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
                for colidx,(cell,value) in enumerate(zip(row.cells,values)):
                    cell.width=Pt(widths[colidx]);cell._tc.remove(cell.paragraphs[0]._p)
                    tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
                    for side,v in [('top',60),('bottom',60),('left',60),('right',60)]:
                        el=OxmlElement('w:'+side);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');marg.append(el)
                    tcpr.append(marg)
                    for val in value if isinstance(value,list) else [value]:
                        if isinstance(val,Paragraph):p=paragraph(cell,val,cell=True)
                        else:p=cell.add_paragraph(str(val))
                        p.paragraph_format.keep_with_next=False
                        if rowidx==0:
                            for r in p.runs:r.bold=True
            for col,w in zip(table.columns,widths):col.width=Pt(w)
            borders=OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            table._tbl.tblPr.append(borders)
            after=doc.add_paragraph();after.paragraph_format.space_after=Pt(3);after.paragraph_format.space_before=Pt(0)
            after.paragraph_format.line_spacing=1;after.add_run().font.size=Pt(2)
        else:raise TypeError(type(item))
    doc.save(path)
    assert Path(path).stat().st_size<1_000_000

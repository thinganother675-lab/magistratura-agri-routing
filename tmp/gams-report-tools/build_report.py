from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base=Path(r'C:/sarychev/Codex/magistratura/coursework/gams-2026/lab2-variant8')
out=base/'report'; out.mkdir(exist_ok=True)
d=Document(); texts=[]
for st in ['Normal','Title','Heading 1','Heading 2']:
 s=d.styles[st]; s.font.name='Times New Roman'; s.font.color.rgb=RGBColor(0,0,0)
 for el in list(s.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
d.styles['Normal'].font.size=Pt(12)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.08
for sn,size in [('Title',18),('Heading 1',14),('Heading 2',12)]: d.styles[sn].font.size=Pt(size)
def setup(s,land=False):
 s.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
 s.page_width=Cm(29.7 if land else 21); s.page_height=Cm(21 if land else 29.7)
 s.top_margin=s.bottom_margin=Cm(1.7); s.left_margin=s.right_margin=Cm(1.8)
setup(d.sections[0])
def p(t,style=None):
 texts.append(t); return d.add_paragraph(t,style)
def h(t): return p(t,'Heading 1')
def table(headers,rows,widths=None,size=11):
 texts.append(' | '.join(headers)); texts.extend(' | '.join(map(str,row)) for row in rows)
 t=d.add_table(rows=1, cols=len(headers)); t.autofit=False
 for c,tx in zip(t.rows[0].cells,headers): c.text=tx
 for row in rows:
  for c,tx in zip(t.add_row().cells,row): c.text=str(tx)
 for ir,row in enumerate(t.rows):
  for ic,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[ic])
   tcpr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   tcpr.append(borders)
   if ir==0:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E6E6E6'); tcpr.append(sh)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(4); par.paragraph_format.space_before=Pt(4)
    par.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(headers)>5 or ic else WD_ALIGN_PARAGRAPH.LEFT
    for r in par.runs:r.font.size=Pt(size); r.bold=ir==0
  trpr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); trpr.append(no)
 return t
p('Лабораторная работа 2 пункт 2','Title')
p('Линейная производственная задача в GAMS')
p('Вариант 8')
p('Выполнил Сарычев Антоний Евгеньевич\nГруппа ДЭ39-26')
h('Цель работы')
p('Построить линейную модель суточного выпуска молочной продукции и найти план с максимальной выручкой. Оптимальный план предусматривает выпуск кефира, сметаны и обезжиренного кефира при полном использовании сырья и общей линии.')
h('Исходные данные')
p('Контракт на сырое молоко уже оплачен. Поэтому целевая функция выражает выручку без повторного вычитания стоимости сырья. Выпуск измеряется в центнерах за сутки; 1 ц = 100 кг.')
table(['Показатель','Молоко x1','Кефир x2','Сметана x3','Обезжиренный кефир x4'],[
 ['Цена, тыс. руб./ц','2.4','2.7','13.8','0.790'],
 ['Расход сырья, ц/ц','1.01','1.01','9.45','0.31'],
 ['Производительность, ц/ч','5','6','0.3','6']],[4.2,3,3,2.8,4.3],size=10)
p('Запас сырья составляет 140 ц/сут. Молоко и оба вида кефира производятся на общей линии, доступной 21 ч/сут. Отдельная линия сметаны доступна 16 ч/сут; её производительность 30 кг/ч = 0.3 ц/ч. Доля обезжиренного кефира в суммарном выпуске двух видов кефира должна быть не менее 1/3. Цена 790 руб./ц равна 0.790 тыс. руб./ц.')
h('Математическая модель')
p('x1, x2, x3, x4 ≥ 0: суточный выпуск молока, кефира, сметаны и обезжиренного кефира соответственно, ц/сут. R: выручка, тыс. руб./сут.')
for s in ['max R = 2.4x1 + 2.7x2 + 13.8x3 + 0.790x4',
'1.01x1 + 1.01x2 + 9.45x3 + 0.31x4 ≤ 140   (сырьё, ц/сут)',
'x1/5 + x2/6 + x4/6 ≤ 21   (общая линия, ч/сут)',
'x3/0.3 ≤ 16   (линия сметаны, ч/сут)',
'x4 ≥ (x2 + x4)/3, или x2 − 2x4 ≤ 0   (ассортимент, ц/сут)']:p(s)
p('Линейная форма условия ассортимента допускает нулевой выпуск обоих видов кефира и не требует деления на их суммарный выпуск.')
d.add_page_break()
h('Решение в GAMS')
p('Модель решена как задача линейного программирования решателем CPLEX. В листинге GAMS получены Solver Status = 1 (Normal Completion) и Model Status = 1 (Optimal).')
table(['Продукт','Выпуск, ц/сут'],[['Молоко','0'],['Кефир','84'],['Сметана','4.4593'],['Обезжиренный кефир','42']],[10,6])
p('Максимальная выручка R = 321.52 тыс. руб./сут. Неокруглённый расчёт: x3 = 602/135 = 4.459259259… ц/сут, R = 321.517777777… тыс. руб./сут. Далее ограничения проверены по неокруглённому плану (0; 84; 602/135; 42).')
h('Проверка ограничений')
table(['Ресурс или условие','Использовано','Предел или требование','Остаток'],[
 ['Сырьё, ц/сут','140','140','0'],['Общая линия, ч/сут','21','21','0'],['Линия сметаны, ч/сут','14.864','16','1.1358'],['Доля обезжиренного кефира','42/(84 + 42) = 1/3','не менее 1/3','0']],[5.4,4,4,3.5],10)
for s in ['Сырьё: 1.01·84 + 9.45·(602/135) + 0.31·42 = 84.84 + 42.14 + 13.02 = 140.',
'Общая линия: 84/6 + 42/6 = 14 + 7 = 21 ч.',
'Линия сметаны: (602/135)/0.3 = 1204/81 = 14.864197531… ≤ 16 ч.',
'Ассортимент: 84 − 2·42 = 0. Все компоненты плана неотрицательны.',
'Выручка: 2.7·84 + 13.8·(602/135) + 0.790·42 = 321.517777777… тыс. руб./сут.']:p(s)
h('Экономический вывод')
p('При данных ценах и мощностях наибольшую выручку даёт выпуск 84 ц кефира, около 4.4593 ц сметаны и 42 ц обезжиренного кефира в сутки. Пастеризованное молоко в оптимальный план не входит. Сырьё и общая линия используются полностью. На линии сметаны остаётся около 1.1358 ч; увеличение только её доступного времени не повысит выручку при остальных неизменных условиях. Требуемая доля обезжиренного кефира достигается точно.')
d.add_page_break()
h('Исходный код GAMS')
p('Множество j задаёт продукты milk, kefir, cream, skim. Переменная x(j) соответствует x1, x2, x3, x4; revenue соответствует R. Множество r задаёт четыре ограничения; price(j), capacity(r) и a(r,j) содержат цены, правые части и коэффициенты.')
p('GAMS создаёт четыре переменные выпуска из x(j) и четыре ограничения из resource(r). Суммирование sum(j, ...) формирует целевую функцию и левые части ограничений. Для расширения ассортимента или перечня ресурсов достаточно дополнить множества и данные без переписывания objective и resource.')
code=(base/'lab2.gms').read_text(encoding='utf-8').strip()
texts.append(code)
for line in code.splitlines():
 par=d.add_paragraph(); par.paragraph_format.space_after=Pt(2)
 r=par.add_run(line); r.font.name='Consolas'; r.font.size=Pt(9)
d.add_page_break()
h('Построение симплексных таблиц')
p('Для канонической формы введены остаточные переменные: sraw для сырья, ц/сут; sline для общей линии, ч/сут; scream для линии сметаны, ч/сут; sassort = 2·skim − kefir, ц/сут. Они неотрицательны. rhs обозначает правую часть; столбец «Базис» указывает базисную переменную. Небазисные переменные равны нулю.')
p('В строке obj используется запись R + Σ(коэффициент·переменная) = rhs. Отрицательная оценка допускает увеличение выручки. Начальный базис образуют четыре остаточные переменные. Все правые части неотрицательны, поэтому искусственные переменные не нужны. Начальная точка вырождена: sassort = 0.')
p('Использовано правило Бланда: входит первая переменная с отрицательной оценкой в порядке milk, kefir, cream, skim, sraw, sline, scream, sassort. Выходит переменная по минимальному неотрицательному отношению rhs/a при a > 0; при равенстве выбирается базисная переменная с меньшим номером. Разрешающая строка делится на разрешающий элемент, затем входящий столбец обнуляется в остальных строках.')
p('Таблицы построены отдельной реализацией прямого симплексного метода в GAMS. Они показывают полный учебный ход решения; последовательность внутренних итераций CPLEX может отличаться. В приложении показаны 6 знаков после точки; вычисления выполнены без промежуточного округления.')
src=(base/'simplex_tables.txt').read_text()
notes=[
'Начальный план равен нулю, R = 0. Входит milk, выходит sline. Отношения: сырьё 138.613861386; общая линия 105. Разрешающий элемент 0.2, шаг 105.',
'Получено milk = 105, R = 252. Входит kefir, выходит sassort. Отношения: сырьё 201.683168317; общая линия 126; ассортимент 0. Разрешающий элемент 1, шаг 0. Это вырожденный переход: меняется базис, план и выручка остаются прежними.',
'Входит cream, выходит sraw. Отношения: сырьё 3.592592593; линия сметаны 4.8. Разрешающий элемент 9.45, шаг 3.592592593.',
'Получено milk = 105, cream = 3.592592593, R = 301.577777778. Входит skim, выходит milk. Отношения: общая линия 42; линия сметаны 58.512820513. Разрешающий элемент 2.5, шаг 42.',
'Все оценки в строке obj неотрицательны. При неотрицательных переменных выручка не может превысить rhs = 321.517777778. План достигает этой величины, поэтому он оптимален. Базис: cream, skim, scream, kefir. Небазисные milk, sraw, sline, sassort равны нулю.']
for i in range(5):
 s=d.add_section(WD_SECTION_START.NEW_PAGE); setup(s,True)
 h(f'Приложение Симплексная таблица {i}')
 p(notes[i])
 block=re.search(r'TABLE '+str(i)+r'[^\n]*\nNonbasic:([^\n]*)\nrow[^\n]*\n((?:[^\n]+\n){5})',src)
 p('Небазисные переменные: '+', '.join(block.group(1).split())+'.')
 rows=[]
 for ln in block.group(2).splitlines():
  vals=ln.split(); rows.append([vals[1]]+vals[2:])
 table(['Базис','milk','kefir','cream','skim','sraw','sline','scream','sassort','rhs'],rows,[2.1]+[2.4]*9,10)
 p('Последняя строка соответствует obj; в ней в столбце «Базис» указана R (z). Единицы rhs: ц/сут для объёмов продукции и сырья, ч/сут для времени, тыс. руб./сут для выручки.')
d.core_properties.author='Сарычев Антоний Евгеньевич'
d.core_properties.title='Лабораторная работа 2 пункт 2 Вариант 8'
d.save(out/'lab2-variant8-sarychev.docx')
(out/'report.txt').write_text('\n\n'.join(texts),encoding='utf-8')
print(out/'report.txt')
